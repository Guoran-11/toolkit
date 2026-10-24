"""核心层单元测试：安全原语、存储工具、转换矩阵。

这些是最容易出现"看起来能跑，边界上出错"的地方，因此重点覆盖边界与攻击面：
    · 路径穿越   —— ``assert_within`` / ``safe_filename``
    · 口令与令牌 —— 哈希往返、签名防篡改、过期
    · 类型嗅探   —— magic bytes 判定
    · 矩阵一致性 —— 每条声明支持的路径都必须有实现分派
"""

from __future__ import annotations

import time
from pathlib import Path

import pytest

# ---------------------------------------------------------------- 路径安全


def test_assert_within_accepts_child(tmp_path: Path) -> None:
    """根目录内的子路径应被放行。"""
    from app.core.security import assert_within

    root = tmp_path / "root"
    root.mkdir()
    child = root / "sub" / "file.txt"
    child.parent.mkdir()

    assert assert_within(child, root) == child.resolve()


@pytest.mark.parametrize(
    "relative",
    ["../escape.txt", "../../etc/passwd", "sub/../../escape.txt"],
)
def test_assert_within_rejects_escape(tmp_path: Path, relative: str) -> None:
    """逃出根目录的路径必须抛错（防御路径穿越）。"""
    from app.core.errors import ForbiddenError
    from app.core.security import assert_within

    root = tmp_path / "root"
    root.mkdir()

    with pytest.raises(ForbiddenError):
        assert_within(root / relative, root)


def test_is_safe_relative() -> None:
    """相对名判定：拒绝分隔符、上级引用、NUL。"""
    from app.core.security import is_safe_relative

    assert is_safe_relative("report.pdf")
    assert is_safe_relative("a b c.xlsx")

    assert not is_safe_relative("")
    assert not is_safe_relative("..")
    assert not is_safe_relative(".")
    assert not is_safe_relative("a/b")
    assert not is_safe_relative("a\\b")
    assert not is_safe_relative("bad\x00name")


# ---------------------------------------------------------------- 文件名净化


@pytest.mark.parametrize(
    "raw",
    [
        "../../etc/passwd",
        "..\\..\\windows\\system32",
        'weird:"<>|?*name.docx',
        "tab\tname.pdf",
        "   ",
    ],
)
def test_safe_filename_strips_dangerous_chars(raw: str) -> None:
    """净化后的文件名不得包含分隔符与保留字符，且不得为空。"""
    from app.core.storage import safe_filename

    cleaned = safe_filename(raw)

    assert cleaned, f"净化结果不应为空：{raw!r}"
    assert "/" not in cleaned
    assert "\\" not in cleaned
    for ch in ':"<>|?*':
        assert ch not in cleaned


def test_safe_filename_preserves_cjk() -> None:
    """中文文件名必须保留（离线单机版的核心使用场景）。"""
    from app.core.storage import safe_filename

    assert "报告" in safe_filename("年度报告.docx")


# ---------------------------------------------------------------- 口令


def test_password_hash_roundtrip() -> None:
    """正确口令通过，错误口令失败；同一口令两次哈希不同（含盐）。"""
    from app.core.security import hash_password, verify_password

    stored = hash_password("S3cret-密码!")
    assert verify_password("S3cret-密码!", stored)
    assert not verify_password("wrong", stored)
    assert hash_password("same") != hash_password("same")


def test_verify_password_handles_malformed_hash() -> None:
    """损坏的哈希串不能让验证崩溃，应安全地返回 False。"""
    from app.core.security import verify_password

    assert verify_password("x", "") is False
    assert verify_password("x", "not-a-valid-hash") is False
    assert verify_password("x", "pbkdf2:onlythree") is False


# ---------------------------------------------------------------- 令牌


def test_token_issue_and_verify() -> None:
    """签发的令牌可被验证，且携带正确的用户名。"""
    from app.core.security import TokenManager

    manager = TokenManager(ttl=3600)
    issued = manager.issue("admin")

    assert "token" in issued
    assert issued["user"]["username"] == "admin"

    payload = manager.verify(issued["token"])
    assert payload["sub"] == "admin"


def test_token_tamper_is_rejected() -> None:
    """篡改签名或载荷必须被拒绝（防伪造）。"""
    from app.core.errors import UnauthorizedError
    from app.core.security import TokenManager

    manager = TokenManager(ttl=3600)
    token = manager.issue("admin")["token"]

    # 改最后一字符 → 签名不匹配
    tampered = token[:-1] + ("A" if token[-1] != "A" else "B")
    with pytest.raises(UnauthorizedError):
        manager.verify(tampered)

    with pytest.raises(UnauthorizedError):
        manager.verify("garbage.token.value")


def test_token_expiry_is_enforced() -> None:
    """过期令牌必须被拒绝。"""
    from app.core.errors import UnauthorizedError
    from app.core.security import TokenManager

    manager = TokenManager(ttl=-1)  # 立即过期
    token = manager.issue("admin")["token"]
    time.sleep(0.01)
    with pytest.raises(UnauthorizedError):
        manager.verify(token)


# ---------------------------------------------------------------- 类型嗅探


def test_sniff_type_recognizes_common_magic() -> None:
    """magic bytes 判定应识别 PDF / ZIP 容器 / 图片。"""
    from app.core.storage import sniff_type

    assert sniff_type(b"%PDF-1.7\n...") == "pdf"
    # DOCX/XLSX/PPTX 都是 ZIP 容器
    assert sniff_type(b"PK\x03\x04rest") == "zip"
    # JPEG 在不同实现中可能返回 jpg 或 jpeg，两者都算正确
    assert sniff_type(b"\xff\xd8\xff\xe0") in {"jpg", "jpeg"}
    assert sniff_type(b"\x89PNG\r\n\x1a\n") == "png"


def test_assert_format_allowed_rejects_mismatch(tmp_path: Path) -> None:
    """声明为 PDF 但内容不是 PDF 时必须拒绝（防扩展名伪装）。

    注意：本函数抛出的具体类型是 ``UnsupportedFormatError``（``ToolkitError``
    的子类），它是 4xx 语义的业务异常，而非底层 ValidationError。
    """
    from app.core.errors import ToolkitError
    from app.core.storage import assert_format_allowed

    fake = tmp_path / "fake.pdf"
    fake.write_bytes(b"this is plainly text, not a pdf")

    with pytest.raises(ToolkitError):
        assert_format_allowed(fake, allowed_kinds=("pdf",), allowed_suffixes=(".pdf",))


# ---------------------------------------------------------------- 审计脱敏


def test_redact_command_masks_secrets() -> None:
    """审计日志中的敏感参数必须被打码（不落明文口令）。"""
    from app.core.security import redact_command

    masked = redact_command("mysql -u root -p SuperSecret123 --execute 'select 1'")
    assert "SuperSecret123" not in masked


# ---------------------------------------------------------------- 转换矩阵


def test_matrix_lookup_known_pair() -> None:
    """已知的转换路径应被识别为受支持。"""
    from app.modules.converter.matrix import lookup

    cell = lookup("docx", "pdf")
    assert cell is not None


def test_matrix_rejects_unknown_pair() -> None:
    """矩阵外的格式对返回 None（由路由层翻译成 400/422）。"""
    from app.modules.converter.matrix import lookup

    assert lookup("docx", "docx") is None or lookup("docx", "docx").grade is None
    assert lookup("mp3", "pdf") is None
    assert lookup("", "") is None


def test_matrix_api_payload_is_serializable() -> None:
    """矩阵的对外载荷必须可 JSON 序列化（否则接口 500）。"""
    import json

    from app.modules.converter.matrix import matrix_api_payload

    payload = matrix_api_payload()
    json.dumps(payload)  # 不抛异常即通过


def test_supported_targets_consistent_with_matrix() -> None:
    """``supported_targets`` 的返回必须与逐格查询一致（防止两处定义漂移）。"""
    from app.modules.converter.matrix import lookup, supported_targets

    for source in ("docx", "xlsx", "pptx", "pdf", "html"):
        targets = supported_targets(source)
        for target in targets:
            cell = lookup(source, target)
            assert cell is not None, f"{source}→{target} 被声明支持但矩阵中查不到"


# ---------------------------------------------------------------- 输出登记


def test_file_registry_register_and_get(tmp_path: Path) -> None:
    """登记产物后可按 file_id 取回，且元信息完整。"""
    from app.core.storage import FileRegistry

    registry = FileRegistry()

    src = tmp_path / "out.pdf"
    src.write_bytes(b"%PDF-1.7 fake-but-registrable")

    record = registry.register(src, name="结果.pdf", mime="application/pdf", ttl_hours=1)

    assert record.file_id
    assert record.name == "结果.pdf"
    assert record.size > 0

    fetched = registry.get(record.file_id)
    assert fetched.file_id == record.file_id
    # url 是 to_api() 派生的对外字段（前端读 output.url）
    assert fetched.to_api()["url"] == f"/api/download/{record.file_id}"


def test_file_registry_rejects_unsafe_id() -> None:
    """登记表的 get 不得接受带路径分隔符的 id。"""
    from app.core.errors import NotFoundError
    from app.core.storage import FileRegistry

    registry = FileRegistry()
    with pytest.raises(NotFoundError):
        registry.get("../secret")


# ---------------------------------------------------------------- 任务产出暂存


def test_conversion_job_stages_output_left_in_workdir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """处理器把产出留在任务目录里时，任务层必须把它搬出去再登记。

    回归背景：``_office_to_pdf`` 曾经漏调 ``_stage()``，产出先被 ``rmtree(workdir)``
    删掉，随后登记抛 ``FileNotFoundError``，用户看到的是"引擎缺失"这种与真实
    原因无关的报错（Word→PDF 直接不可用）。
    """
    import shutil

    from app.modules.converter import service
    from app.modules.converter.base import ConversionOutcome, ConversionParams

    workdir = tmp_path / "job-staging"
    workdir.mkdir()
    source = tmp_path / "报告.docx"
    source.write_bytes(b"PK\x03\x04fake-docx")

    def fake_run(request) -> ConversionOutcome:
        produced = request.workdir / "output.pdf"
        produced.write_bytes(b"%PDF-1.7 staging-regression")
        return ConversionOutcome(path=produced, name="报告.pdf", mime="application/pdf")

    class _Ctx:
        """只实现 job 用到的 progress 接口。"""

        def progress(self, _value: int) -> None:
            pass

    monkeypatch.setattr(service, "run_conversion", fake_run)
    monkeypatch.setattr(service.ConversionRequest, "create_workdir", classmethod(lambda cls: workdir))

    job = service.make_conversion_job(
        source_path=source,
        source_format="docx",
        target_format="pdf",
        display_name="报告.docx",
        params=ConversionParams(),
    )

    try:
        result = job(_Ctx())
        assert result.path.exists(), "产出被任务目录回收逻辑误删（漏调 _stage 的典型症状）"
        assert result.path.parent != workdir, "产出仍留在任务目录内"
        assert not workdir.exists(), "任务目录未被回收"
    finally:
        # 暂存区不归本次测试管理，手工清理避免磁盘残留（NFR-03）
        for leftover in (tmp_path / "job-staging",):
            shutil.rmtree(leftover, ignore_errors=True)
        from app.core.config import TEMP_DIR

        for staged in TEMP_DIR.glob("staged-*.pdf"):
            staged.unlink(missing_ok=True)


def test_workdir_handlers_stage_their_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """在任务目录里产出文件的处理器，返回值必须已搬出该目录。

    覆盖历史上漏调 ``_stage()`` 的五条路径：
        docx/xlsx/pptx → pdf、html → pdf、html → png、html → docx、html → xlsx。
    引擎函数全部被替换为"只写一个文件"的替身，因此本测试不依赖 LibreOffice/Chromium。
    """
    import shutil

    from app.modules.converter import office, service
    from app.modules.converter.base import ConversionParams, ConversionRequest

    class _Ctx:
        """只实现 progress 接口。"""

        def progress(self, _value: int) -> None:
            pass

    produced: list[Path] = []

    def _writer(suffix: str):
        def _fake(source: Path, workdir: Path, *args, **kwargs) -> Path:
            workdir.mkdir(parents=True, exist_ok=True)
            target = workdir / f"output.{suffix}"
            target.write_bytes(b"fake-output")
            produced.append(target)
            return target

        return _fake

    monkeypatch.setattr(service, "office_to_pdf", _writer("pdf"))
    monkeypatch.setattr(service, "render_html_to_pdf", _writer("pdf"))
    monkeypatch.setattr(service, "render_html_to_image", _writer("png"))
    monkeypatch.setattr(office, "html_to_docx", _writer("docx"))
    monkeypatch.setattr(office, "html_to_xlsx", _writer("xlsx"))

    # _require_office_docx 要求文件头是 OOXML（zip）而非 OLE 旧格式
    ooxml = tmp_path / "文档.docx"
    ooxml.write_bytes(b"PK\x03\x04fake-docx")
    plain = tmp_path / "页面.html"
    plain.write_text("<html><body>hi</body></html>", encoding="utf-8")

    cases = [
        ("office→pdf", service._office_to_pdf, ooxml, "docx", "pdf"),
        ("html→pdf", service._html_to_pdf, plain, "html", "pdf"),
        ("html→png", service._html_to_png, plain, "html", "png"),
        ("html→docx", service._html_to_docx, plain, "html", "docx"),
        ("html→xlsx", service._html_to_xlsx, plain, "html", "xlsx"),
    ]

    staged_paths: list[Path] = []
    try:
        for label, handler, source, source_format, target_format in cases:
            workdir = tmp_path / f"job-{label.replace('→', '-')}"
            workdir.mkdir()
            request = ConversionRequest(
                source_path=source,
                source_format=source_format,
                target_format=target_format,
                params=ConversionParams(),
                display_name=source.name,
                workdir=workdir,
                ctx=_Ctx(),
            )
            outcome = handler(request)
            staged_paths.append(outcome.path)
            assert outcome.path.exists(), f"{label}：产出不存在（被任务目录回收误删）"
            # 独立判断归属，不复用被测代码的 _inside（否则改动它会让断言一起失真）
            assert workdir.resolve() not in outcome.path.resolve().parents, f"{label}：产出仍留在任务目录内"
    finally:
        for path in produced + staged_paths:
            path.unlink(missing_ok=True)
        for directory in tmp_path.glob("job-*"):
            shutil.rmtree(directory, ignore_errors=True)


# ---------------------------------------------------------------- XLSX 兜底导出

def test_write_tables_to_xlsx_falls_back_to_text_rows(tmp_path: Path) -> None:
    """没有任何表格时，正文文本必须仍能导出为工作表，而不是直接失败。"""
    pytest.importorskip("openpyxl")

    import openpyxl

    from app.modules.converter.documents import _write_tables_to_xlsx

    output = _write_tables_to_xlsx(
        [],
        tmp_path / "out.xlsx",
        one_sheet_per_table=True,
        source_label="文档",
        fallback_rows=[["第一段正文"], ["   "], ["第二段正文"]],
        fallback_sheet_title="正文",
    )

    workbook = openpyxl.load_workbook(output)
    assert workbook.sheetnames == ["正文"]
    values = [row[0] for row in workbook["正文"].iter_rows(values_only=True)]
    assert values == ["第一段正文", "第二段正文"]


def test_write_tables_to_xlsx_keeps_tables_and_text(tmp_path: Path) -> None:
    """有表格时表格独立成表，正文另存一表，两者都不丢。"""
    pytest.importorskip("openpyxl")

    import openpyxl

    from app.modules.converter.documents import _write_tables_to_xlsx

    output = _write_tables_to_xlsx(
        [[["姓名", "分数"], ["张三", "90"]]],
        tmp_path / "out.xlsx",
        one_sheet_per_table=True,
        source_label="文档",
        fallback_rows=[["表头说明"]],
    )

    workbook = openpyxl.load_workbook(output)
    assert {"表格1", "正文"} <= set(workbook.sheetnames)
    assert workbook["表格1"]["A1"].value == "姓名"
    assert workbook["正文"]["A1"].value == "表头说明"


def test_write_tables_to_xlsx_requires_any_content(tmp_path: Path) -> None:
    """表格与文本都为空时才报错，且错误文案指向 PDF 兜底方案。"""
    from app.core.errors import ConversionFailedError
    from app.modules.converter.documents import _write_tables_to_xlsx

    with pytest.raises(ConversionFailedError) as exc_info:
        _write_tables_to_xlsx(
            [],
            tmp_path / "out.xlsx",
            one_sheet_per_table=True,
            source_label="文档",
            fallback_rows=[[], ["  "]],
        )
    assert "转 PDF" in str(exc_info.value)
